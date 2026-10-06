from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor

out = Path(__file__).resolve().parent
doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.27)
sec.page_height = Inches(11.69)
sec.top_margin = sec.bottom_margin = Inches(0.7)
sec.left_margin = sec.right_margin = Inches(0.75)
for name in ['Normal', 'Title', 'Heading 1']:
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style.font.color.rgb = RGBColor(0, 0, 0)
normal = doc.styles['Normal']
normal.font.size = Pt(11)
normal.paragraph_format.line_spacing = 1.08
normal.paragraph_format.space_after = Pt(7)
doc.styles['Title'].font.size = Pt(21)
doc.styles['Title'].paragraph_format.space_after = Pt(5)
doc.styles['Heading 1'].font.size = Pt(12)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(8)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(4)

title = 'EURUSD Release Sequence Analysis'
meta = 'Report date: 6 October 2026 | Release window: 3-12 August 2026'
intro = ('Overall assessment: Mildly bullish EUR/USD after all three release groups. '
         'The initial dollar support from ISM weakens after payroll and wage deterioration. '
         'Monthly CPI acceleration then limits the EUR/USD upside case.')
method = ('Method: Actual minus previous (A-P) deltas only. Payroll comparisons use the displayed '
          'A-P column consistently. Index changes are in points; rate changes are in percentage points (pp).')
sections = [
('ISM manufacturing and services', 'Release dates: 3 and 5 August 2026',
 'Manufacturing PMI (+2.3), employment (+3.1) and new orders (+0.7) strengthen, while prices paid ease (-1.9). '
 'This initially favors USD and weighs on EUR/USD. Services activity (+3.7), orders (+2.1) and prices paid (+2.6) '
 'reinforce demand and inflation pressure, but employment deteriorates (-3.8). The combined ISM signal is mildly bearish EUR/USD, with a labor-market warning.'),
('Nonfarm payrolls', 'Release date: 7 August 2026',
 'Payroll momentum deteriorates by 80k, comprising government (-61k) and private payrolls (-19k). '
 'Government accounts for about 76% of the deterioration, limiting evidence of broad private-sector weakness. '
 'Wage growth slows by 0.2 pp monthly and 0.3 pp annually. Unemployment falls 0.1 pp, but participation also falls 0.1 pp, '
 'making the unemployment improvement less convincing; these deltas do not establish its cause. Hours and U6 are unchanged. '
 'The balance favors USD weakness and shifts the EUR/USD bias bullish.'),
('Consumer price inflation', 'Release date: 12 August 2026',
 'Monthly headline inflation accelerates by 0.5 pp and core inflation by 0.2 pp. However, headline and core annual inflation '
 'each slow by 0.1 pp. Giving monthly changes greater weight for the latest inflation direction, CPI is a mildly dollar-positive '
 'counterweight. It supports a possible EUR/USD pullback and reduces the bullish conviction established after NFP.'),
('Final interpretation', None,
 'Bias progression: mildly bearish after ISM, bullish after NFP, then mildly bullish after CPI. '
 'Weaker employment and wages tilt the sequence toward easier Fed policy, while renewed monthly inflation pressure limits that case. '
 'This is a qualitative US-data bias, assuming unchanged euro-area conditions; the deltas do not establish an exchange-rate target or observed price reaction.')]

doc.add_paragraph(title, 'Title')
p = doc.add_paragraph(meta)
p.runs[0].font.size = Pt(10)
p = doc.add_paragraph(intro)
p.runs[0].bold = True
p = doc.add_paragraph(method)
p.runs[0].font.size = Pt(10)
for heading, date, body in sections:
    doc.add_paragraph(heading, 'Heading 1')
    p = doc.add_paragraph()
    if date:
        r = p.add_run(date + '. ')
        r.bold = True
    p.add_run(body)
doc.core_properties.title = title
doc.core_properties.subject = 'ISM NFP and CPI implications for EURUSD using actual minus previous deltas'
doc.core_properties.author = ''
path = out / 'EURUSD_Release_Report_2026-10-06.docx'
doc.save(path)
content = [title, meta, intro, method]
for heading, date, body in sections:
    content.extend([heading, (date + '. ' if date else '') + body])
print(path)
print('Word count:', len(' '.join(content).split()))
loaded = Document(path)
assert len(loaded.paragraphs) == 12
assert all(s in '\n'.join(p.text for p in loaded.paragraphs) for s in ['6 October 2026', '7 August 2026', '12 August 2026', '0.5 pp', '80k'])
print('Content and document structure verified')
